"""Deterministic CV extraction and prompt-injection screening.

CV text is untrusted data. This module removes visually hidden text from the
text passed to any AI service and flags instruction-like content for human
review. It never asks an LLM to decide whether a document is safe.
"""

from __future__ import annotations

from dataclasses import dataclass
import hashlib
import io
import re
from typing import Iterable

import fitz
from docx import Document
from docx.oxml.ns import qn


INJECTION_PATTERNS = (
    re.compile(r"\bignore\s+(?:all\s+)?(?:previous|prior|above)\s+instructions?\b", re.I),
    re.compile(r"\b(?:system|developer|assistant)\s+(?:prompt|message|instruction)s?\b", re.I),
    re.compile(r"\b(?:give|assign|award|set)\s+(?:me|this candidate|the candidate)?\s*(?:a\s*)?(?:score|rating|marks?)\b", re.I),
    re.compile(r"\b(?:shortlist|hire|select|recommend)\s+(?:me|this candidate|the candidate)\b", re.I),
    re.compile(r"\b(?:override|bypass|disregard)\s+(?:the\s+)?(?:rules|criteria|instructions?|screening)\b", re.I),
    re.compile(r"\b(?:do not|don't)\s+(?:evaluate|reject|screen)\b", re.I),
    re.compile(r"\byou\s+are\s+(?:an?\s+)?(?:ai|language model|recruiting bot)\b", re.I),
)


@dataclass(frozen=True)
class CVSecurityResult:
    safe_text: str
    status: str
    flags: tuple[str, ...]
    sha256: str
    extracted_characters: int

    @property
    def requires_manual_review(self) -> bool:
        return self.status != "CLEAN"

    def public_metadata(self) -> dict:
        return {
            "status": self.status,
            "flags": list(self.flags),
            "sha256": self.sha256,
            "scanner_version": "1.0.0",
            "extracted_characters": self.extracted_characters,
        }


def _instruction_found(text: str) -> bool:
    return any(pattern.search(text) for pattern in INJECTION_PATTERNS)


def _near_white(color: int | None) -> bool:
    if color is None:
        return False
    red, green, blue = (color >> 16) & 255, (color >> 8) & 255, color & 255
    return red >= 242 and green >= 242 and blue >= 242


def _scan_pdf(content: bytes) -> tuple[list[str], list[str], set[str]]:
    visible: list[str] = []
    hidden: list[str] = []
    flags: set[str] = set()
    with fitz.open(stream=content, filetype="pdf") as document:
        if document.page_count > 100:
            raise ValueError("CV exceeds the 100-page safety limit")
        for page in document:
            data = page.get_text("dict")
            for block in data.get("blocks", []):
                for line in block.get("lines", []):
                    for span in line.get("spans", []):
                        text = str(span.get("text", "")).strip()
                        if not text:
                            continue
                        bbox = span.get("bbox", (0, 0, 0, 0))
                        outside_page = bbox[2] <= 0 or bbox[3] <= 0 or bbox[0] >= page.rect.width or bbox[1] >= page.rect.height
                        suspicious_style = _near_white(span.get("color")) or float(span.get("size", 12)) < 2 or outside_page
                        if suspicious_style:
                            hidden.append(text)
                            flags.add("HIDDEN_OR_INVISIBLE_TEXT")
                        else:
                            visible.append(text)
    return visible, hidden, flags


def _docx_run_hidden(run) -> bool:
    properties = run._element.get_or_add_rPr()
    vanished = properties.find(qn("w:vanish")) is not None
    color = run.font.color.rgb
    near_white = bool(color and color[0] >= 242 and color[1] >= 242 and color[2] >= 242)
    tiny = bool(run.font.size and run.font.size.pt < 2)
    return vanished or near_white or tiny


def _scan_docx(content: bytes) -> tuple[list[str], list[str], set[str]]:
    document = Document(io.BytesIO(content))
    visible: list[str] = []
    hidden: list[str] = []
    flags: set[str] = set()
    paragraphs: Iterable = list(document.paragraphs) + [p for table in document.tables for row in table.rows for cell in row.cells for p in cell.paragraphs]
    for paragraph in paragraphs:
        for run in paragraph.runs:
            text = run.text.strip()
            if not text:
                continue
            if _docx_run_hidden(run):
                hidden.append(text)
                flags.add("HIDDEN_OR_INVISIBLE_TEXT")
            else:
                visible.append(text)
    return visible, hidden, flags


def scan_cv(content: bytes, content_type: str) -> CVSecurityResult:
    digest = hashlib.sha256(content).hexdigest()
    try:
        if content_type == "application/pdf":
            visible, hidden, flags = _scan_pdf(content)
        elif content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            visible, hidden, flags = _scan_docx(content)
        else:
            raise ValueError("Unsupported CV format")
    except Exception:
        return CVSecurityResult("", "UNSCANNABLE", ("TEXT_EXTRACTION_FAILED",), digest, 0)

    safe_text = "\n".join(visible).strip()
    visible_injection = _instruction_found(safe_text)
    hidden_injection = _instruction_found("\n".join(hidden))
    if visible_injection:
        flags.add("PROMPT_INJECTION_TEXT")
    if hidden_injection:
        flags.add("HIDDEN_PROMPT_INJECTION")
    if not safe_text:
        flags.add("NO_EXTRACTABLE_VISIBLE_TEXT")
    status = "SUSPICIOUS" if flags else "CLEAN"
    return CVSecurityResult(safe_text, status, tuple(sorted(flags)), digest, len(safe_text))
